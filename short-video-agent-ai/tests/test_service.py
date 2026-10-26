import copy
import json
import math
from pathlib import Path

import httpx
import pytest
from fastapi.testclient import TestClient

from app import demo
from app.config import Settings
from app.errors import ServiceError
from app.knowledge import chunk_text, extract_file
from app.main import create_app
from app.models import GenerateRequest


HEADERS = {"X-Internal-Key": "test-internal-secret"}


def request(phase="GENERATE", **extra):
    return {"taskId": "task-1", "phase": phase,
            "task": {"topic": "大学生防诈骗", "audience": "大学生", "platform": "抖音", "style": "轻松", "duration": 60,
                     "requirements": "有据可查", "category": "科普", "ragEnabled": True, "mode": "QUICK"}, **extra}


class MockUpstreams:
    """Real HTTP wire schemas (httpx transport), with a small in-memory Qdrant."""
    def __init__(self):
        self.points = {}
        self.dimension = None
        self.calls = []
        self.fail_model = False
        self.invalid_model = 0
        self.wrong_duration = False

    @staticmethod
    def matches(point, filter):
        payload = point["payload"]
        def condition(item):
            return (payload.get(item["key"]) in item["match"]["any"] if "any" in item["match"]
                    else payload.get(item["key"]) == item["match"]["value"])
        return all(condition(item) for item in filter.get("must", [])) and not any(condition(item) for item in filter.get("must_not", []))

    def __call__(self, req):
        body = json.loads(req.content) if req.content else {}
        self.calls.append((req.method, req.url.path, body))
        if req.url.host == "models.test":
            assert req.headers["authorization"] == "Bearer gen-secret"
            if self.fail_model:
                return httpx.Response(401, json={"error": "credentials must not be reflected"})
            if self.invalid_model:
                self.invalid_model -= 1
                return httpx.Response(200, json={"choices": [{"message": {"content": "not-json"}}]})
            schema = body["messages"][0]["content"]
            if '"brief"' in schema:
                result = {"brief": "防诈科普", "query": "大学生防诈骗"}
            elif '"topics"' in schema:
                result = {"topics": [topic.model_dump() for topic in demo.topics(GenerateRequest.model_validate(request()))]}
            elif '"usedChunkIds"' in schema:
                value = demo.content(GenerateRequest.model_validate(request())).model_dump()
                value["content"]["qualityNotes"] = ["Mocked real provider output"]
                if self.wrong_duration and "校验警告" not in body["messages"][-1]["content"]:
                    for segment in value["content"]["segments"]:
                        segment["duration"] = 1
                available = list(self.points.values())
                value["usedChunkIds"] = [available[0]["id"]] if available else []
                result = value
            else:
                result = {"positioning": "大学生防诈科普", "audienceAnalysis": "易受网络诱导的大学生", "outline": ["场景", "建议", "互动"]}
            return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps(result, ensure_ascii=False)}}],
                                           "usage": {"prompt_tokens": 20, "completion_tokens": 40}})
        if req.url.host == "embedding.test":
            assert req.headers["authorization"] == "Bearer embed-secret"
            from app.providers import demo_embedding
            return httpx.Response(200, json={"data": [{"index": i, "embedding": demo_embedding(text, 16)} for i, text in enumerate(body["input"])]})
        path = req.url.path
        if "/collections/" in path and "/points" not in path:
            if req.method == "PUT":
                self.dimension = body["vectors"]["size"]
                return httpx.Response(200, json={"result": True})
            return httpx.Response(200 if self.dimension else 404, json={"result": {"config": {"params": {"vectors": {"size": self.dimension}}}}})
        if path.endswith("/points"):
            for point in body["points"]:
                self.points[point["id"]] = copy.deepcopy(point)
        elif path.endswith("/points/delete"):
            self.points = {key: p for key, p in self.points.items() if not self.matches(p, body["filter"])}
        elif path.endswith("/points/payload"):
            for point in self.points.values():
                if self.matches(point, body["filter"]):
                    point["payload"].update(body["payload"])
        elif path.endswith("/points/query"):
            matches = []
            for point in self.points.values():
                if self.matches(point, body["filter"]):
                    a, b = body["query"], point["vector"]
                    score = sum(x * y for x, y in zip(a, b)) / (math.sqrt(sum(x*x for x in a)) * math.sqrt(sum(y*y for y in b)))
                    if score >= body["score_threshold"]:
                        matches.append({"payload": point["payload"], "score": score})
            return httpx.Response(200, json={"result": {"points": sorted(matches, key=lambda p: p["score"], reverse=True)[:body["limit"]]}})
        else:
            return httpx.Response(404, json={})
        return httpx.Response(200, json={"result": True})


@pytest.fixture
def upstream():
    return MockUpstreams()


def client(upstream, mode="demo", **kwargs):
    settings = Settings(mode=mode, internal_key="test-internal-secret", qdrant_url="http://vectors.test",
                        generation_base_url="http://models.test/v1", generation_api_key="gen-secret", generation_model="test-model",
                        embedding_base_url="http://embedding.test/v1", embedding_api_key="embed-secret", embedding_model="test-embedding", **kwargs)
    return TestClient(create_app(settings, transport=httpx.MockTransport(upstream)))


def index(cli, text="大学生防诈骗：遇到网络转账需要先核实身份，保护个人信息，不向陌生人提供验证码。", **extra):
    return cli.post("/internal/knowledge/index", headers=HEADERS,
                    json={"documentId": "doc-1", "name": "防诈骗资料", "text": text, "category": "科普", "sourceUrl": "https://example.com/source", "enabled": True, **extra})


def test_private_auth_and_no_secret_disclosure(upstream):
    with client(upstream) as cli:
        assert cli.get("/health").json() == {"status": "ok"}
        assert cli.get("/internal/health").status_code == 401
        assert cli.get("/internal/health", headers={"X-Internal-Key": "wrong"}).status_code == 401
        response = cli.get("/internal/health", headers=HEADERS)
        assert response.status_code == 200
        assert "secret" not in response.text
        assert response.json()["statusDetail"].startswith("DEMO")


def test_default_missing_real_config_returns_503(upstream):
    with TestClient(create_app(Settings(internal_key="test-internal-secret"), transport=httpx.MockTransport(upstream))) as cli:
        response = cli.post("/internal/generate", headers=HEADERS, json=request())
        assert response.status_code == 503
        assert response.json()["code"] == "GENERATION_NOT_CONFIGURED"
        response = index(cli)
        assert response.status_code == 503
        assert response.json()["code"] == "EMBEDDING_NOT_CONFIGURED"
    with TestClient(create_app(Settings())) as cli:
        assert cli.get("/internal/health", headers=HEADERS).status_code == 503


def test_plan_and_generation_demo(upstream):
    with client(upstream) as cli:
        assert index(cli).status_code == 200
        plan = cli.post("/internal/generate", headers=HEADERS, json=request("PLAN"))
        assert plan.status_code == 200, plan.text
        assert len(plan.json()["topics"]) == 3
        assert "content" not in plan.json()
        assert len(plan.json()["sources"]) == 1
        result = cli.post("/internal/generate", headers=HEADERS, json=request(selectedTopic=plan.json()["topics"][1]))
        assert result.status_code == 200, result.text
        data = result.json()
        assert sum(segment["duration"] for segment in data["content"]["segments"]) == 60
        assert data["quality"]["mode"] == "demo"
        assert len(data["sources"]) == 1
        assert "[资料:" + data["sources"][0]["chunkId"] + "]" in data["content"]["segments"][1]["narration"]
        assert data["sources"][0]["text"][:80] in data["content"]["segments"][1]["narration"]
        assert data["usage"]["model"] == "DEMO-FIXTURE"


@pytest.mark.parametrize("section,allowed", [("title", {"titles"}), ("opening", {"opening"}), ("script", {"segments"}),
                                               ("ending", {"ending"}), ("caption", {"caption", "tags"}), ("1", {"segments"})])
def test_revision_scope_exact_and_preserve_deleted_source_snapshot(upstream, section, allowed):
    with client(upstream) as cli:
        index(cli)
        plan = cli.post("/internal/generate", headers=HEADERS, json=request("PLAN")).json()
        content = demo.content(GenerateRequest.model_validate(request())).content.model_dump()
        cli.delete("/internal/knowledge/doc-1", headers=HEADERS)
        response = cli.post("/internal/generate", headers=HEADERS,
                            json=request("REVISE", content=content, instruction="更直接一点", section=section, sources=plan["sources"]))
        assert response.status_code == 200, response.text
        data = response.json()
        for key in content:
            if key not in allowed:
                assert data["content"][key] == content[key]
        if section.isdigit():
            for i in (0, 2, 3):
                assert data["content"]["segments"][i] == content["segments"][i]
            assert data["content"]["segments"][1] != content["segments"][1]
        assert data["sources"] == plan["sources"]


def test_vector_lifecycle_replacement_and_filters(upstream):
    with client(upstream) as cli:
        first = index(cli).json()
        assert first["chunkCount"] == 1
        assert index(cli).json() == first
        assert len(upstream.points) == 1
        search = lambda category="科普": cli.post("/internal/knowledge/search", headers=HEADERS, json={"query": "大学生防诈骗", "category": category, "limit": 5})
        assert len(search().json()) == 1
        assert search("美食").json() == []
        assert cli.put("/internal/knowledge/doc-1", headers=HEADERS, json={"enabled": False, "category": "科普", "sourceUrl": ""}).status_code == 200
        assert search().json() == []
        second = index(cli, text="大学生防诈骗：谨防刷单诈骗，保护验证码，不向未知账户转账。").json()
        assert second["chunks"][0]["id"] != first["chunks"][0]["id"]
        assert len(upstream.points) == 1
        assert search().json()[0]["chunkId"] == second["chunks"][0]["id"]
        assert cli.delete("/internal/knowledge/doc-1", headers=HEADERS).status_code == 200
        assert search().json() == []


def test_real_model_and_embedding_http_integration(upstream):
    upstream.wrong_duration = True
    with client(upstream, "real") as cli:
        assert index(cli).status_code == 200
        result = cli.post("/internal/generate", headers=HEADERS, json=request())
        assert result.status_code == 200, result.text
        data = result.json()
        assert data["quality"]["mode"] == "real"
        assert data["quality"]["repairs"] == 1
        assert len(data["sources"]) == 1
        assert data["usage"]["inputTokens"] == 100
        assert data["usage"]["outputTokens"] == 200
        assert data["usage"]["model"] == "test-model"
        assert any(event["stage"] == "REPAIR" for event in data["events"])
        assert any(path.endswith("embeddings") for _, path, _ in upstream.calls)


def test_model_invalid_json_bounded_retry_and_auth_failure(upstream):
    upstream.invalid_model = 3
    with client(upstream, "real") as cli:
        response = cli.post("/internal/generate", headers=HEADERS, json=request("PLAN"))
        assert response.status_code == 502
        assert response.json()["code"] == "INVALID_MODEL_OUTPUT"
        assert len(upstream.calls) == 2
        upstream.fail_model = True
        response = cli.post("/internal/generate", headers=HEADERS, json=request("PLAN"))
        assert response.status_code == 502
        assert "credentials must" not in response.text


def test_file_path_and_empty_text(tmp_path):
    root = tmp_path / "knowledge"
    root.mkdir()
    outside = tmp_path / "private.txt"
    outside.write_text("private")
    with pytest.raises(ServiceError):
        extract_file(str(outside), root)
    (root / "link.txt").symlink_to(outside)
    with pytest.raises(ServiceError):
        extract_file(str(root / "link.txt"), root)
    valid = root / "test.md"
    valid.write_text("# 大学生防诈骗\n核实身份。")
    assert "大学生" in extract_file(str(valid), root)
    with pytest.raises(ServiceError):
        chunk_text("  \n ")
    text = "测试文本。" * 500
    chunks = chunk_text(text)
    assert len(chunks) > 1
    assert all(0 < len(chunk) <= 700 for chunk in chunks)


def test_request_validation_and_invalid_revision(upstream):
    with client(upstream) as cli:
        body = request("REVISE", instruction="secret request must not reflect", section="99")
        response = cli.post("/internal/generate", headers=HEADERS, json=body)
        assert response.status_code == 422
        assert "secret request" not in response.text
        response = cli.post("/internal/knowledge/index", headers=HEADERS,
                            json={"documentId": "d", "name": "x", "text": "a", "path": "/etc/passwd"})
        assert response.status_code == 422


def test_docx_paragraphs_tables_and_empty_pdf(tmp_path):
    from docx import Document
    from pypdf import PdfWriter
    doc = Document()
    doc.add_paragraph("短视频创作资料")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "口播"
    table.cell(0, 1).text = "画面"
    path = tmp_path / "knowledge.docx"
    doc.save(path)
    extracted = extract_file(str(path), tmp_path)
    assert "短视频创作资料" in extracted
    assert "口播 | 画面" in extracted
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    pdf = tmp_path / "empty.pdf"
    writer.write(pdf)
    with pytest.raises(ServiceError) as exc:
        chunk_text(extract_file(str(pdf), tmp_path))
    assert exc.value.code == "EMPTY_DOCUMENT"


def test_low_score_no_match(upstream):
    with client(upstream, retrieval_min_score=0.99) as cli:
        index(cli)
        response = cli.post("/internal/knowledge/search", headers=HEADERS, json={"query": "今天煮什么饭"})
        assert response.status_code == 200
        assert response.json() == []


@pytest.mark.parametrize("alteration", ["blank_title", "string_duration", "unknown_field", "citation"])
def test_malformed_model_output_never_accepted(upstream, alteration):
    def handler(req):
        response = upstream(req)
        if req.url.path.endswith("chat/completions") and response.status_code == 200:
            data = response.json()
            value = json.loads(data["choices"][0]["message"]["content"])
            if "content" in value:
                if alteration == "blank_title":
                    value["content"]["titles"] = ["   "]
                elif alteration == "string_duration":
                    value["content"]["segments"][0]["duration"] = "15"
                elif alteration == "unknown_field":
                    value["content"]["untrustedExtra"] = "arbitrary"
                else:
                    value["content"]["opening"] += "[资料:invented-chunk]"
                    value["usedChunkIds"] = ["invented-chunk"]
                data["choices"][0]["message"]["content"] = json.dumps(value)
                return httpx.Response(200, json=data)
        return response
    cfg = Settings(mode="real", internal_key="test-internal-secret", generation_api_key="gen-secret",
                   generation_base_url="http://models.test/v1", generation_model="model",
                   qdrant_url="http://vectors.test")
    with TestClient(create_app(cfg, transport=httpx.MockTransport(handler))) as cli:
        body = request()
        body["task"]["ragEnabled"] = False
        response = cli.post("/internal/generate", headers=HEADERS, json=body)
        assert response.status_code == 502
        assert response.json()["code"] == ("INVALID_CITATIONS" if alteration == "citation" else "INVALID_MODEL_OUTPUT")


def test_private_library_scope_is_applied_before_vector_top_k(upstream):
    with client(upstream, mode="real", retrieval_min_score=-1) as cli:
        assert index(cli, documentId="public").status_code == 200
        assert index(cli, documentId="alice", ownerId="user-alice").status_code == 200
        assert index(cli, documentId="bob", ownerId="user-bob").status_code == 200
        body = {"query": "大学生防诈骗", "limit": 1}
        response = cli.post("/internal/knowledge/search", headers=HEADERS, json=body)
        assert [s["documentId"] for s in response.json()] == ["public"]
        response = cli.post("/internal/knowledge/search", headers=HEADERS,
                            json={**body, "documentIds": ["alice"]})
        assert [s["documentId"] for s in response.json()] == ["alice"]
        request_filter = upstream.calls[-1][2]["filter"]
        assert {"key": "documentId", "match": {"any": ["alice"]}} in request_filter["must"]
        calls = len(upstream.calls)
        response = cli.post("/internal/knowledge/search", headers=HEADERS,
                            json={**body, "documentIds": []})
        assert response.json() == []
        assert len(upstream.calls) == calls


def test_chunk_preview_uses_the_same_settings_as_index_without_embedding(upstream):
    with TestClient(create_app(Settings(internal_key="test-internal-secret"), transport=httpx.MockTransport(upstream))) as cli:
        text = ("拍摄前先确定主体和光线。\n" * 60)
        response = cli.post("/internal/knowledge/preview", headers=HEADERS,
                            json={"text": text, "chunkSize": 300, "chunkOverlap": 60})
        assert response.status_code == 200
        preview = response.json()
        assert preview["chunkCount"] == len(chunk_text(text, 300, 60))
        assert [c["text"] for c in preview["chunks"]] == chunk_text(text, 300, 60)
        assert upstream.calls == []
        response = cli.post("/internal/knowledge/preview", headers=HEADERS,
                            json={"text": text, "chunkSize": 200, "chunkOverlap": 200})
        assert response.status_code == 422
    with client(upstream) as cli:
        first = index(cli, text="拍摄前先确定主体和光线。" * 100, chunkSize=300, chunkOverlap=50).json()
        second = index(cli, text="拍摄前先确定主体和光线。" * 100, chunkSize=600, chunkOverlap=100).json()
        assert first["chunkCount"] > second["chunkCount"]
        assert len(upstream.points) == second["chunkCount"]
