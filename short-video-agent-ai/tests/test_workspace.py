import asyncio
import zipfile
from pathlib import Path
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import NameObject,DictionaryObject,DecodedStreamObject
import httpx
import pytest
from app.documents import extract_document
from app.errors import ServiceError
from app.config import Settings
from app.conversation import Context,ConversationAgent
from app.conversation_models import Brief,ConversationRequest
from app.providers import Provider


def test_document_formats_and_upload_boundary(tmp_path):
    doc=Document();doc.add_paragraph('校园创作计划');doc.add_table(rows=1,cols=1).cell(0,0).text='手机拍摄';doc.save(tmp_path/'sample.docx')
    assert '手机拍摄' in extract_document(str(tmp_path/'sample.docx'),tmp_path)['text']
    with zipfile.ZipFile(tmp_path/'sample.pptx','w') as archive:
        archive.writestr('ppt/presentation.xml','<p:presentation xmlns:p="p" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:sldId r:id="first"/><p:sldId r:id="second"/></p:presentation>')
        archive.writestr('ppt/_rels/presentation.xml.rels','<Relationships><Relationship Id="first" Type="x/slide" Target="slides/slide1.xml"/><Relationship Id="second" Type="x/slide" Target="slides/slide2.xml"/></Relationships>')
        archive.writestr('ppt/slides/slide2.xml','<p:sld xmlns:p="p" xmlns:a="a"><a:t>第二页讲剪辑</a:t></p:sld>')
        archive.writestr('ppt/slides/slide1.xml','<p:sld xmlns:p="p" xmlns:a="a"><a:t>第一页讲拍摄</a:t></p:sld>')
    result=extract_document(str(tmp_path/'sample.pptx'),tmp_path)['text']
    assert result.index('第一页') < result.index('第二页')
    with zipfile.ZipFile(tmp_path/'sample.odt','w') as archive:
        archive.writestr('content.xml','<body xmlns:text="text"><text:p>创作目标与受众</text:p></body>')
    assert '创作目标' in extract_document(str(tmp_path/'sample.odt'),tmp_path)['text']
    with pytest.raises(ServiceError,match='路径无效'):
        extract_document(str(tmp_path.parent/'not-allowed.txt'),tmp_path)
    (tmp_path/'bad.pdf').write_text('not a pdf')
    with pytest.raises(ServiceError,match='无法解析'):
        extract_document(str(tmp_path/'bad.pdf'),tmp_path)


def test_text_pdf_and_scanned_pdf(tmp_path):
    writer=PdfWriter();page=writer.add_blank_page(width=300,height=300)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
    stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 20 200 Td (Creative brief: mobile filming) Tj ET')
    page[NameObject('/Contents')]=writer._add_object(stream)
    writer.write(tmp_path/'text.pdf')
    assert 'mobile filming' in extract_document(str(tmp_path/'text.pdf'),tmp_path)['text']
    blank=PdfWriter();blank.add_blank_page(width=100,height=100);blank.write(tmp_path/'scan.pdf')
    with pytest.raises(ServiceError,match='OCR'):
        extract_document(str(tmp_path/'scan.pdf'),tmp_path)


@pytest.mark.asyncio
async def test_document_paging_and_history_access_are_scoped_to_request():
    request=ConversationRequest(conversationId='one',runId='run',messages=[{'role':'user','content':'读文件与历史作品'}],documents=[{'id':'file-one','name':'简报.docx','text':'a'*900+'结尾'}],historyWorks=[{'id':'work-one','title':'咖啡日记','summary':'手冲咖啡','content':{'opening':'旧稿开场'}}])
    cfg=Settings(mode='demo')
    async with httpx.AsyncClient() as client:
        agent=ConversationAgent(cfg,Provider(cfg,client),None)
        ctx=Context(request,Brief(),None,[],[])
        result,_=await agent.execute(ctx,'read_document',{'documentId':'file-one','offset':500,'limit':500})
        assert result['text'].endswith('结尾') and not result['hasMore']
        found,_=await agent.execute(ctx,'search_history_creations',{'query':'咖啡'})
        assert found['count']==1 and 'content' not in found['works'][0]
        read,_=await agent.execute(ctx,'read_history_creation',{'artifactId':'work-one'})
        assert read['content']['opening']=='旧稿开场'
        with pytest.raises(ServiceError,match='没有找到'):
            await agent.execute(ctx,'read_history_creation',{'artifactId':'other-user-work'})
        with pytest.raises(ServiceError,match='没有这个文档'):
            await agent.execute(ctx,'read_document',{'documentId':'other-user-file'})
