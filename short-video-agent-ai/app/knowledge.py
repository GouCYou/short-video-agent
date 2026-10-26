import asyncio
import hashlib
import math
import uuid
import zipfile
from collections import defaultdict
from pathlib import Path
from urllib.parse import quote
from docx import Document
from pypdf import PdfReader
import httpx
from .config import Settings
from .errors import ServiceError
from .models import IndexRequest, MetadataRequest, SearchRequest, Source
from .providers import Provider, request_json


def extract_file(raw_path: str, root: Path) -> str:
    """Only regular files within the configured shared volume; never fetch URLs."""
    path = Path(raw_path).resolve()
    if not path.is_relative_to(root.resolve()) or not path.is_file():
        raise ServiceError("知识文件必须位于 KNOWLEDGE_DIR 内且为普通文件。", 400, "INVALID_FILE")
    if path.stat().st_size > 20 * 1024 * 1024:
        raise ServiceError("文件超过 20 MB 限制。", 413, "FILE_TOO_LARGE")
    try:
        if path.suffix.lower() in {".txt", ".md"}:
            data = path.read_bytes()
            try:
                return data.decode("utf-8-sig")
            except UnicodeDecodeError:
                return data.decode("gb18030")
        if path.suffix.lower() == ".docx":
            with zipfile.ZipFile(path) as archive:
                if sum(info.file_size for info in archive.infolist()) > 80 * 1024 * 1024:
                    raise ServiceError("DOCX 解压后超过限制。", 413, "FILE_TOO_LARGE")
            document = Document(path)
            parts = [p.text for p in document.paragraphs]
            parts.extend(" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
            return "\n".join(parts)
        if path.suffix.lower() == ".pdf":
            reader = PdfReader(path)
            if reader.is_encrypted:
                raise ServiceError("暂不支持加密 PDF，请上传可提取文本的文件。", 400, "INVALID_FILE")
            if len(reader.pages) > 500:
                raise ServiceError("PDF 页数超过 500 页限制。", 413, "FILE_TOO_LARGE")
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        raise ServiceError("仅支持 TXT、Markdown、DOCX、文本型 PDF。", 400, "INVALID_FILE")
    except ServiceError:
        raise
    except Exception as exc:
        raise ServiceError("无法解析文件，请检查文件格式与内容。", 400, "PARSE_ERROR") from exc


def chunk_text(text: str, size: int = 700, overlap: int = 100) -> list[str]:
    text = text.replace("\x00", "").replace("\r\n", "\n").strip()
    if not text:
        raise ServiceError("未提取到文本；扫描 PDF 需要先进行 OCR。", 400, "EMPTY_DOCUMENT")
    if len(text) > 2_000_000:
        raise ServiceError("提取文本超过 200 万字符限制。", 413, "FILE_TOO_LARGE")
    if not 0 <= overlap < size:
        raise ValueError("invalid chunk overlap")
    result, start = [], 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            boundary = max(text.rfind("\n", start + size // 2, end), text.rfind("。", start + size // 2, end))
            if boundary >= 0:
                end = boundary + 1
        value = text[start:end].strip()
        if value:
            result.append(value)
        if end == len(text):
            break
        start = max(start + 1, end - overlap)
    return result


class Knowledge:
    def __init__(self, settings: Settings, provider: Provider, client: httpx.AsyncClient):
        self.cfg, self.provider, self.client = settings, provider, client
        self._locks: dict[str, asyncio.Lock] = defaultdict(asyncio.Lock)
        self._collection_lock = asyncio.Lock()

    @property
    def base(self):
        return self.cfg.qdrant_url + "/collections/" + quote(self.cfg.collection, safe="")

    @property
    def headers(self):
        return {"api-key": self.cfg.qdrant_api_key} if self.cfg.qdrant_api_key else {}

    async def call(self, method: str, suffix: str, **kwargs):
        return await request_json(self.client, method, self.base + suffix, headers=self.headers,
                                  retries=self.cfg.retries, **kwargs)

    async def exists(self):
        try:
            result = await self.client.get(self.base, headers=self.headers)
        except httpx.HTTPError as exc:
            raise ServiceError("Qdrant 无法连接，请检查 QDRANT_URL。", 503, "VECTOR_UNAVAILABLE") from exc
        if result.status_code == 404:
            return None
        if result.status_code != 200:
            raise ServiceError("Qdrant 访问异常，请检查地址和凭据。", 503, "VECTOR_UNAVAILABLE")
        try:
            return result.json()["result"]
        except (ValueError, KeyError) as exc:
            raise ServiceError("Qdrant 返回异常。") from exc

    async def ensure_collection(self, dimension: int):
        async with self._collection_lock:
            info = await self.exists()
            if info:
                existing = info["config"]["params"]["vectors"]["size"]
                if existing != dimension:
                    raise ServiceError("Embedding 维度与已有知识库不匹配。更换 QDRANT_COLLECTION 并重新索引。", 409, "EMBEDDING_DIMENSION_MISMATCH")
            else:
                await self.call("PUT", "", json={"vectors": {"size": dimension, "distance": "Cosine"}})

    @staticmethod
    def document_filter(document_id: str):
        return {"must": [{"key": "documentId", "match": {"value": document_id}}]}

    async def index(self, req: IndexRequest):
        text = req.text if req.text is not None else await asyncio.to_thread(extract_file, req.path, self.cfg.knowledge_dir)
        chunks = chunk_text(text, req.chunkSize, req.chunkOverlap)
        # Embedding before deleting the old revision keeps prior data on API failures.
        vectors = []
        for offset in range(0, len(chunks), 16):
            vectors.extend(await self.provider.embed(chunks[offset:offset + 16]))
        await self.ensure_collection(len(vectors[0]))
        revision = hashlib.sha256((text + self.cfg.mode + self.cfg.embedding_model + f":{req.chunkSize}:{req.chunkOverlap}").encode()).hexdigest()
        rows = [{"id": str(uuid.uuid5(uuid.NAMESPACE_URL, req.documentId + ":" + revision + ":" + str(i))), "text": chunk}
                for i, chunk in enumerate(chunks)]
        async with self._locks[req.documentId]:
            points = [{"id": row["id"], "vector": vectors[i], "payload": {
                "documentId": req.documentId, "documentName": req.name, "text": row["text"],
                "chunkId": row["id"], "category": req.category, "sourceUrl": req.sourceUrl,
                "enabled": req.enabled, "revision": revision,
                "visibility": "private" if req.ownerId else "public",
            }} for i, row in enumerate(rows)]
            for offset in range(0, len(points), 64):
                await self.call("PUT", "/points?wait=true", json={"points": points[offset:offset + 64]})
            stale = self.document_filter(req.documentId)
            stale["must_not"] = [{"key": "revision", "match": {"value": revision}}]
            await self.call("POST", "/points/delete?wait=true", json={"filter": stale})
        return {"chunkCount": len(rows), "chunks": rows}

    async def search(self, req: SearchRequest, document_ids: list[str] | None = None) -> list[Source]:
        if document_ids == []:
            return []
        # Missing embedding credentials are explicit, including on an empty collection.
        vector = (await self.provider.embed([req.query]))[0]
        info = await self.exists()
        if not info:
            return []
        if info["config"]["params"]["vectors"]["size"] != len(vector):
            raise ServiceError("Embedding 维度与知识库不一致，请重新索引。", 409, "EMBEDDING_DIMENSION_MISMATCH")
        must = [{"key": "enabled", "match": {"value": True}}]
        if req.category:
            must.append({"key": "category", "match": {"value": req.category}})
        scope = {"must": must}
        if document_ids is None:
            scope["must_not"] = [{"key": "visibility", "match": {"value": "private"}}]
        else:
            must.append({"key": "documentId", "match": {"any": document_ids}})
        result = await self.call("POST", "/points/query", json={"query": vector, "filter": scope,
                                "limit": req.limit, "with_payload": True, "score_threshold": self.cfg.retrieval_min_score})
        try:
            sources, seen = [], set()
            for row in result["result"]["points"]:
                payload = row["payload"]
                if payload["text"] in seen or not payload.get("enabled") or not math.isfinite(row["score"]):
                    continue
                seen.add(payload["text"])
                sources.append(Source(documentId=payload["documentId"], documentName=payload["documentName"],
                                      chunkId=payload["chunkId"], text=payload["text"], score=row["score"], sourceUrl=payload.get("sourceUrl", "")))
            return sources
        except (KeyError, ValueError, TypeError) as exc:
            raise ServiceError("知识检索返回结构异常。") from exc

    async def update(self, document_id: str, req: MetadataRequest):
        async with self._locks[document_id]:
            if await self.exists():
                await self.call("POST", "/points/payload?wait=true", json={"payload": req.model_dump(), "filter": self.document_filter(document_id)})
        return {"ok": True}

    async def delete(self, document_id: str):
        async with self._locks[document_id]:
            if await self.exists():
                await self.call("POST", "/points/delete?wait=true", json={"filter": self.document_filter(document_id)})
        return {"ok": True}
